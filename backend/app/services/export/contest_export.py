"""参赛图文科普导出增强 — AI 使用声明、赛制命名、合规提示"""
from __future__ import annotations

import io
from typing import Any


def build_ai_declaration_text(ai_declaration: dict | None) -> str:
    """根据 AI 声明配置生成声明文本"""
    if not ai_declaration:
        return ""

    parts = ["AI 辅助创作声明"]
    if ai_declaration.get("ai_image"):
        tools = ai_declaration.get("image_tools", "")
        parts.append(f"- 本作品部分配图使用 AI 工具辅助生成" + (f"（{tools}）" if tools else ""))
    if ai_declaration.get("ai_text"):
        parts.append("- 本作品部分文字内容使用 AI 工具辅助撰写")
    if ai_declaration.get("human_reviewed"):
        parts.append("- 所有内容均经人工审校")

    if len(parts) <= 1:
        return ""
    return "\n".join(parts)


def get_contest_filename(article: Any, contest_pack: Any = None) -> str:
    """生成符合赛制要求的文件名"""
    naming_template = None

    if contest_pack and hasattr(contest_pack, 'naming_template') and contest_pack.naming_template:
        naming_template = contest_pack.naming_template

    custom_rules = getattr(article, 'contest_custom_rules', None)
    if custom_rules and isinstance(custom_rules, dict) and custom_rules.get('naming_template'):
        naming_template = custom_rules['naming_template']

    if naming_template:
        title = (article.title or article.topic or "作品名").replace("/", "-").replace("\\", "-")
        info = getattr(article, 'submission_info', None) or {}
        result = naming_template
        result = result.replace("作品名", title)
        result = result.replace("单位", info.get("unit", "单位"))
        result = result.replace("科室", info.get("department", "科室"))
        result = result.replace("第一作者", info.get("author", "作者"))
        return result

    title = (article.topic or "参赛作品").replace("/", "-").replace("\\", "-")
    return title


def get_contest_font_config(article: Any, contest_pack: Any = None) -> dict | None:
    """获取赛制要求的字体配置"""
    font = None

    if contest_pack and hasattr(contest_pack, 'font') and contest_pack.font:
        font = contest_pack.font

    custom_rules = getattr(article, 'contest_custom_rules', None)
    if custom_rules and isinstance(custom_rules, dict) and custom_rules.get('font'):
        font = custom_rules['font']

    if not font:
        return None

    import re
    size_match = re.search(r'(\d+)\s*号', font)
    font_name = re.sub(r'\d+\s*号', '', font).strip()

    size_map = {
        '一': 26, '二': 22, '三': 16, '四': 14, '小四': 12,
        '五': 10.5, '1': 26, '2': 22, '3': 16, '4': 14, '5': 10.5,
    }
    pt_size = None
    if size_match:
        pt_size = size_map.get(size_match.group(1))

    return {
        "font_name": font_name or "仿宋",
        "pt_size": pt_size,
    }


def append_ai_declaration_to_docx(doc: Any, ai_declaration: dict | None) -> None:
    """在 DOCX 文档末尾追加 AI 使用声明"""
    text = build_ai_declaration_text(ai_declaration)
    if not text:
        return

    from docx.shared import Pt, RGBColor

    doc.add_paragraph()  # blank line
    for line in text.split("\n"):
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        if line == "AI 辅助创作声明":
            run.bold = True
            run.font.size = Pt(10)


def get_confirmation_message(article: Any, contest_pack: Any = None) -> str | None:
    """生成导出前的确认提示文本"""
    if contest_pack:
        pack_name = getattr(contest_pack, 'name', '')
        source_note = getattr(contest_pack, 'source_note', '')
        updated_at = getattr(contest_pack, 'updated_at', '')

        msg = f"本赛制包「{pack_name}」"
        if source_note:
            msg += f"，来源说明：{source_note}"
        if updated_at:
            msg += f"，最后更新于 {str(updated_at)[:10]}"
        msg += "。请在投递前对照官方最新通知核对，本工具不对赛事方临时调整负责。"
        return msg

    rule_source = getattr(article, 'contest_rule_source', None)
    if rule_source in ('parsed', 'manual'):
        label = '公告解析' if rule_source == 'parsed' else '手工填写'
        return (
            f"当前赛制配置来源：{label}。"
            "请在投递前对照官方最新通知核对字数、格式等要求，本工具不对赛事方临时调整负责。"
        )

    return None
